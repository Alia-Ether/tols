
#│-----------------------------------------------------------│
#│  Link: t.me/FrontendVSCode                 │ 
#│  Author: Frontend & LED (𝙰𝚕𝚒𝚊 𝙴𝚝𝚑𝚎𝚛 𖤍) 🌷 │ 
#│  lang: python                              │
#│  [VS-HASH-01] ΞΩ77Λβ99PPHD8A71             │ 
#│  build:3.10.15                             │ 
#│-----------------------------------------------------------│



from openpyxl import load_workbook, Workbook
from pathlib import Path
from datetime import datetime

def main():
    input_file = Path("sales.xlsx")
    if not input_file.exists():
        wb_empty = Workbook()
        ws_empty = wb_empty.active
        ws_empty.append(["Сумма", "Дата", "Менеджер"])
        wb_empty.save(input_file)
        print("🍎 Файл sales.xlsx не найден, создан новый пустой файл с заголовками!")
        return

    try:
        wb = load_workbook(input_file)
        ws = wb.active
    except Exception as e:
        print("❌ Не удалось открыть Excel:", e)
        return

    headers = [cell.value for cell in ws[1]]
    required_cols = {"Сумма", "Дата", "Менеджер"}
    if not required_cols.issubset(headers):
        print("🍎 Извини, в файле нет нужных колонок:", ", ".join(required_cols))
        return

    idx_sum = headers.index("Сумма")
    idx_date = headers.index("Дата")
    idx_manager = headers.index("Менеджер")

    filtered_rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_date = row[idx_date]
        if isinstance(row_date, str):
            try:
                row_date = datetime.strptime(row_date, "%Y-%m-%d")
            except ValueError:
                continue
        if row[idx_sum] > 10000 and row_date >= datetime(2024, 1, 1):
            filtered_rows.append(list(row) + [row[idx_sum] * 1.2])

    if not filtered_rows:
        print("🍎 Ничего не найдено по условиям фильтрации")
        return

    # filtered_sales.xlsx
    wb_filtered = Workbook()
    ws_filtered = wb_filtered.active
    ws_filtered.append(headers + ["Сумма с НДС"])
    for row in filtered_rows:
        ws_filtered.append(row)
    wb_filtered.save("filtered_sales.xlsx")

    # summary_by_manager.xlsx
    summary = {}
    for row in filtered_rows:
        manager = row[idx_manager]
        summary[manager] = summary.get(manager, 0) + row[-1]

    wb_summary = Workbook()
    ws_summary = wb_summary.active
    ws_summary.append(["Менеджер", "Сумма с НДС"])
    for manager, total in sorted(summary.items(), key=lambda x: x[1], reverse=True):
        ws_summary.append([manager, total])
    wb_summary.save("summary_by_manager.xlsx")

    print("🌸 Готово!")
    print("✅ filtered_sales.xlsx")
    print("✅ summary_by_manager.xlsx")


if __name__ == "__main__":
    main()